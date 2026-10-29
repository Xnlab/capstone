"""Exercise the public deployment using disposable fictional accounts only.

No passwords or access tokens are printed. The reported disposable IDs can be
removed after verification without touching real users.
"""
import io
import json
import secrets
import sys
import uuid
import urllib.error
import urllib.request
from http.cookies import SimpleCookie
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document

BASE = sys.argv[1].rstrip("/") + "/api/v1"
password = secrets.token_urlsafe(24)
suffix = uuid.uuid4().hex[:12]
users = []


def call(method, route, data=None, token=None, raw=None, content_type=None, expected=200):
    headers = {}
    if token:
        headers["Cookie"] = token["cookie"]
        if method not in ("GET", "HEAD", "OPTIONS"):
            headers["X-CSRF-Token"] = token["csrf_token"]
    body = raw
    if data is not None:
        body = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    if content_type:
        headers["Content-Type"] = content_type
    req = urllib.request.Request(BASE + route, data=body, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=60) as res:
            result = res.read()
            assert res.status == expected, (route, res.status)
            parsed = json.loads(result) if "application/json" in res.headers.get("Content-Type", "") else result
            if route in ("/auth/register", "/auth/login"):
                cookies = SimpleCookie(res.headers.get("Set-Cookie", ""))
                session = next((item for name, item in cookies.items() if name.endswith("cdoc_session")), None)
                assert session is not None, "HttpOnly session cookie missing"
                parsed["_session"] = {"cookie": session.key + "=" + session.value, "csrf_token": parsed["csrf_token"]}
            return parsed
    except urllib.error.HTTPError as error:
        if error.code == expected:
            return None
        raise RuntimeError(f"{method} {route}: HTTP {error.code}: {error.read(600).decode(errors='replace')}") from None


def upload(name, data, token, expected=200):
    boundary = "smoke" + uuid.uuid4().hex
    body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()
            + data + f"\r\n--{boundary}--\r\n".encode())
    return call("POST", "/templates/upload", token=token, raw=body,
                content_type="multipart/form-data; boundary=" + boundary, expected=expected)


try:
    assert call("GET", "/health")["cloud_mode"]
    for index in range(2):
        account = call("POST", "/auth/register", {"email": f"smoke-{suffix}-{index}@cloud-smoke.invalid", "password": password, "name": "배포 검증용 임시 사용자"})
        users.append(account)
    token, other = users[0]["_session"], users[1]["_session"]
    logged = call("POST", "/auth/login", {"email": users[0]["user"]["email"], "password": password})
    assert logged["user"]["id"] == users[0]["user"]["id"]
    call("GET", "/contracts", expected=401)
    call("PUT", "/company", {"company_name": "배포검증 가상회사", "representative": "테스트", "address": "가상 주소", "business_number": "123-45-67890"}, token)
    contract = call("POST", "/contracts", {"contract_title": "클라우드 검증 용역", "client_name": "가상기관", "amount": 11000000, "contract_date": "2026-10-29", "start_date": "2026-10-29", "end_date": "2026-11-30", "description": "배포 확인용 가상 계약"}, token)
    cid = contract["id"]
    call("GET", f"/contracts/{cid}", token=other, expected=404)
    templates = call("GET", "/templates", token=token)
    assert len([t for t in templates if t["kind"] == "builtin"]) == 9
    checks = []
    for t in templates:
        if t["kind"] != "builtin":
            continue
        doc = call("POST", f"/contracts/{cid}/documents", {"template_id": t["id"]}, token)
        did = doc["id"]
        call("GET", f"/documents/{did}", token=other, expected=404)
        native = call("GET", f"/documents/{did}/download?format=native", token=token)
        assert native[:2] == b"PK"
        with ZipFile(io.BytesIO(native)) as archive:
            assert ("xl/workbook.xml" in archive.namelist()) if t["output_format"] == "xlsx" else ("Contents/section0.xml" in archive.namelist())
        pdf = call("GET", f"/documents/{did}/download?format=pdf", token=token)
        with pymupdf.open(stream=pdf, filetype="pdf") as parsed:
            assert len(parsed) >= 1
            assert "클라우드" in "".join(page.get_text() for page in parsed), t["name"]
        preview = call("GET", f"/documents/{did}/preview-pages", token=token)
        assert preview["pages"]
        checks.append(t["name"])
    word = Document(); word.add_paragraph("양식 검증"); word.add_paragraph("상호 : ____")
    buf = io.BytesIO(); word.save(buf)
    uploaded = upload("smoke.docx", buf.getvalue(), token)["template"]
    tid = uploaded["id"]
    call("GET", f"/templates/{tid}", token=other, expected=404)
    assert call("GET", f"/templates/{tid}/original", token=token)[:2] == b"PK"
    fields = uploaded["fields"]
    assert fields
    for f in fields:
        f["field_key"] = "company_name"
    call("PUT", f"/templates/{tid}/mapping", {"fields": fields}, token)
    doc = call("POST", f"/contracts/{cid}/documents", {"template_id": tid}, token)
    native = call("GET", f'/documents/{doc["id"]}/download?format=native', token=token)
    filled = Document(io.BytesIO(native))
    assert "배포검증 가상회사" in "".join(p.text for p in filled.paragraphs)
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "server"))
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "server/tests"))
    from test_hwpx_upload import fixture
    hwpx = upload("smoke.hwpx", fixture(), token)["template"]
    assert hwpx["fields"]
    call("PUT", f'/templates/{hwpx["id"]}/mapping', {"fields": hwpx["fields"]}, token)
    hd = call("POST", f"/contracts/{cid}/documents", {"template_id": hwpx["id"]}, token)
    hn = call("GET", f'/documents/{hd["id"]}/download?format=native', token=token)
    with ZipFile(io.BytesIO(hn)) as archive:
        assert "배포검증 가상회사" in archive.read("Contents/section0.xml").decode()
        assert archive.read("BinData/preserved.png") == b"preserved-original-image"
    call("GET", f'/documents/{hd["id"]}/preview-pages', token=token, expected=503)
    with pymupdf.open() as pdf:
        page = pdf.new_page(); page.insert_text((30, 40), "Company name:")
        raw_pdf = pdf.tobytes()
    pt = upload("smoke.pdf", raw_pdf, token)["template"]
    call("PUT", f'/templates/{pt["id"]}/mapping', {"fields": [{"label": "회사명", "field_key": "company_name", "location": {"type": "pdf_rect", "page": 0, "rect": [30, 70, 400, 100]}}]}, token)
    pd = call("POST", f"/contracts/{cid}/documents", {"template_id": pt["id"]}, token)
    pn = call("GET", f'/documents/{pd["id"]}/download?format=native', token=token)
    with pymupdf.open(stream=pn, filetype="pdf") as pdf:
        assert "배포검증 가상회사" in pdf[0].get_text()
    upload("legacy.hwp", b"unsupported", token, expected=400)
    print(json.dumps({"status": "passed", "builtin_templates": checks, "checks": ["public health", "registration/login", "owner isolation", "all native/PDF/preview outputs", "persistent DOCX/HWPX/PDF upload and autofill", "HWP restriction", "uploaded HWPX PDF restriction"], "disposable_user_ids": [a["user"]["id"] for a in users]}, ensure_ascii=False))
finally:
    if users:
        print(json.dumps({"disposable_user_ids": [a["user"]["id"] for a in users], "email_prefix": "smoke-" + suffix}))
