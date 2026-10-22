"""Professional native document rendering for built-in templates.

HWPX is generated from a sanitized, content-free package scaffold. XLSX uses
real formulas so users can continue editing quantities and unit prices.
"""

import html
import re
from decimal import Decimal, ROUND_HALF_UP
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from docx import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from app.standard_fields import field_label
from app.standard_fields import korean_amount


FONT = "맑은 고딕"
_PH = re.compile(r"\{(\w+)\}")
_DIGITS = re.compile(r"[^0-9-]")
BASE_HWPX = Path(__file__).resolve().parent.parent / "assets" / "base.hwpx"

HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HS = "http://www.hancom.co.kr/hwpml/2011/section"
HWP_CONTENT_WIDTH = 47056  # A4 minus two 22 mm margins, in 1/100 pt.


def item_amounts(item):
    """Whole-won rounding, matching Excel ROUND rather than Python ties-to-even."""
    supply = int((Decimal(str(item["qty"])) * Decimal(item["unit"])).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    tax = int((Decimal(supply) / 10).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return supply, tax


def _fill(text: str, values: dict) -> str:
    return _PH.sub(lambda m: str(values.get(m.group(1), "") or f"[{field_label(m.group(1))}]"), text)


def money(value) -> int:
    digits = _DIGITS.sub("", str(value or ""))
    try:
        return int(digits or 0)
    except ValueError:
        return 0


def parse_items(raw: str, fallback_name: str, fallback_supply: int) -> list[dict]:
    """Parse `item|qty|unit price|note`, one item per line."""
    items = []
    for line in str(raw or "").splitlines():
        if not line.strip():
            continue
        parts = [p.strip() for p in line.split("|")]
        name = parts[0] if parts else ""
        if not name:
            continue
        try:
            qty = float(parts[1]) if len(parts) > 1 and parts[1] else 1
        except ValueError:
            qty = 1
        unit = money(parts[2]) if len(parts) > 2 else 0
        note = parts[3] if len(parts) > 3 else ""
        items.append({"name": name, "qty": qty, "unit": unit, "note": note})
    if not items:
        items.append({"name": fallback_name or "용역 수행", "qty": 1, "unit": fallback_supply, "note": ""})
    return items


def _set_cell_shading(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, *, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(str(text or ""))
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _paragraph(doc, text="", *, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = FONT
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        r.font.size = Pt(size)
        r.bold = bold
    return p


def _doc_base(title: str) -> DocxDocument:
    doc = DocxDocument()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = title.replace(" ", "")
    _paragraph(doc, title, bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    return doc


def _info_table_docx(doc, rows, values, seal=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(3.6), Cm(13)
    for label, key in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(3.6), Cm(13)
        _set_cell_text(cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(cells[0], "F2F3F5")
        _set_cell_text(cells[1], values.get(key, ""))
        if key == 'representative' and seal:
            from app.services.seal import add_docx_seal
            _set_cell_text(cells[1], f"{values.get(key, '')}             (인)", align=WD_ALIGN_PARAGRAPH.RIGHT)
            add_docx_seal(cells[1].paragraphs[0], seal, Cm(13).pt)
            table.rows[-1].height = Cm(seal['height_mm'] / 10 + .4)
    _paragraph(doc, after=5)


def _items_docx(doc, raw, values, mode):
    items = parse_items(raw, values.get("description") or values.get("contract_title", ""), money(values.get("supply_amount")))
    headings = ["항 목", "수량", "단가", "공급가액", "세액", "합계"]
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    for column, width in zip(table.columns, (5.5, 1.1, 2.4, 2.6, 2.2, 2.8)):
        column.width = Cm(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, label in zip(table.rows[0].cells, headings):
        _set_cell_text(c, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(c, "F2F3F5")
    supply_total = tax_total = 0
    for item in items:
        supply, tax = item_amounts(item)
        supply_total += supply
        tax_total += tax
        row = table.add_row().cells
        vals = [item["name"], f"{item['qty']:g}", f"{item['unit']:,}", f"{supply:,}", f"{tax:,}", f"{supply + tax:,}"]
        for idx, (c, value) in enumerate(zip(row, vals)):
            _set_cell_text(c, value, align=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT, size=9.5)
    row = table.add_row().cells
    row[0].merge(row[2])
    _set_cell_text(row[0], "합 계", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(row[3], f"{supply_total:,}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_cell_text(row[4], f"{tax_total:,}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_cell_text(row[5], f"{supply_total + tax_total:,}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for c in row:
        _set_cell_shading(c, "F2F2F2")
    _paragraph(doc, after=5)


def render_docx(definition: dict, values: dict, seal=None) -> DocxDocument:
    values = presentation_values(values)
    doc = _doc_base(definition.get("title", "서류"))
    for section in definition.get("sections", []):
        kind = section.get("type")
        if kind == "page_break":
            doc.add_page_break()
            _paragraph(doc, section["title"], bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
        elif kind == "signature":
            _signature_docx(doc, values, definition.get("recipient"), seal)
        elif kind in ("info", "bank"):
            _info_table_docx(doc, section.get("rows", []), values, seal if not definition.get('signature') else None)
        elif kind == "paragraph":
            _paragraph(doc, _fill(section.get("text", ""), values), size=10.5, after=10)
        elif kind == "field_block":
            table = doc.add_table(rows=2, cols=1)
            table.style = "Table Grid"
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            _set_cell_text(table.cell(0, 0), section["label"], bold=True)
            _set_cell_shading(table.cell(0, 0), "F2F3F5")
            _set_cell_text(table.cell(1, 0), values.get(section["key"], ""), size=10)
            _paragraph(doc, after=5)
        elif kind == "line_items":
            _items_docx(doc, values.get(section.get("key"), ""), values, section.get("mode", "cost"))
        elif kind == "numbered":
            for idx, item in enumerate(section.get("items", []), 1):
                _paragraph(doc, f"{idx}. {_fill(item, values)}", after=5)
        elif kind == "attachments":
            _paragraph(doc, "붙임", bold=True, after=2)
            for idx, item in enumerate(attachment_items(section, values), 1):
                _paragraph(doc, f"  {idx}. {_fill(item, values)}", after=2)
        elif kind in ("stamp_box", "photo_boxes"):
            labels = section.get("labels", [])
            table = doc.add_table(rows=2, cols=max(1, len(labels)))
            table.style = "Table Grid"
            for idx, label in enumerate(labels):
                _set_cell_text(table.cell(0, idx), label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_shading(table.cell(0, idx), "F2F2F2")
                _set_cell_text(table.cell(1, idx), "" if kind == "stamp_box" else "[사진 첨부]", align=WD_ALIGN_PARAGRAPH.CENTER)
            table.rows[1].height = Cm(3.2)
            _paragraph(doc, after=5)

    if definition.get("signature"):
        _signature_docx(doc, values, definition.get("recipient"), seal)
    return doc


def _signature_docx(doc, values, recipient, seal=None):
    _paragraph(doc, values.get("submit_date") or "    년    월    일", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    for label, key in (("주 소", "address"), ("상 호", "company_name"), ("대표자", "representative")):
        suffix = ("             (인)" if seal else "       (인)") if key == "representative" else ""
        p = _paragraph(doc, f"{label} : {values.get(key, '')}{suffix}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
        p.paragraph_format.keep_with_next = True
        if key == 'representative' and seal:
            from app.services.seal import add_docx_seal
            sec = doc.sections[0]
            add_docx_seal(p, seal, sec.page_width.pt - sec.left_margin.pt - sec.right_margin.pt)
    if recipient:
        _paragraph(doc, _fill(recipient, values), bold=True, size=12, after=0)


def presentation_values(values):
    values = dict(values)
    for key, target in (("amount", "contract_amount_display"), ("payment_amount", "claim_amount_display"), ("advance_amount", "advance_amount_display")):
        number = money(values.get(key))
        values[target] = f"{korean_amount(number)} (₩{number:,})" if values.get(key) else ""
    return values


def attachment_items(section, values):
    raw = values.get("attachment_list")
    return [line.strip() for line in raw.splitlines() if line.strip()] if raw is not None else section.get("items", [])


def render_quotation_xlsx(values: dict, seal=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A13"
    ws.merge_cells("A1:F2")
    ws["A1"] = "견 적 서"
    ws["A1"].font = Font(name="맑은 고딕", size=20, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = PatternFill("solid", fgColor="FFF2CC")

    thin = Side(style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    blue, gray, yellow = "DCE6F1", "F2F2F2", "FFF2CC"
    meta = [
        (4, "프로젝트명", values.get("contract_title", ""), "수신자", values.get("client_name", "")),
        (5, "사업자번호", values.get("business_number", ""), "작성일자", values.get("submit_date", "")),
        (6, "상호", values.get("company_name", ""), "대표자", values.get("representative", "") + "  (인)"),
        (7, "업태", values.get("business_type", ""), "종목", values.get("business_item", "")),
        (8, "주소", values.get("address", ""), "연락처", values.get("phone", "")),
    ]
    for row, l1, v1, l2, v2 in meta:
        ws.cell(row, 1, l1); ws.cell(row, 2, v1); ws.cell(row, 4, l2); ws.cell(row, 5, v2)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        for col in range(1, 7):
            c = ws.cell(row, col); c.border = border; c.alignment = Alignment(vertical="center", wrap_text=True)
            c.font = Font(name="맑은 고딕", size=10, bold=col in (1, 4))
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=blue)
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=blue)

    ws.merge_cells("A10:B11"); ws["A10"] = "견적금액 (VAT 포함)"
    ws.merge_cells("C10:F11")
    ws["A10"].fill = PatternFill("solid", fgColor=yellow)
    ws["C10"].fill = PatternFill("solid", fgColor="E2F0D9")
    ws["C10"].number_format = '#,##0"원"'
    for row in range(10, 12):
        for col in range(1, 7):
            ws.cell(row, col).border = border
            ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, col).font = Font(name="맑은 고딕", size=11, bold=True)

    headers = ["품 명", "수량", "단가", "공급가액", "세액", "비고"]
    for col, label in enumerate(headers, 1):
        c = ws.cell(13, col, label); c.fill = PatternFill("solid", fgColor=yellow); c.border = border
        c.font = Font(name="맑은 고딕", size=10, bold=True); c.alignment = Alignment(horizontal="center", vertical="center")

    items = parse_items(values.get("quote_items", ""), values.get("contract_title", ""), money(values.get("supply_amount")))
    start = 14
    for idx, item in enumerate(items, start):
        ws.cell(idx, 1, item["name"]); ws.cell(idx, 2, item["qty"]); ws.cell(idx, 3, item["unit"])
        ws.cell(idx, 4, f"=ROUND(B{idx}*C{idx},0)"); ws.cell(idx, 5, f"=ROUND(D{idx}*0.1,0)"); ws.cell(idx, 6, item["note"])
        for col in range(1, 7):
            c = ws.cell(idx, col); c.border = border; c.font = Font(name="맑은 고딕", size=9.5)
            c.alignment = Alignment(horizontal="left" if col in (1, 6) else "right", vertical="center", wrap_text=True)
        for col in (2, 3, 4, 5):
            ws.cell(idx, col).number_format = "#,##0"

    subtotal = start + len(items)
    total = subtotal + 1
    ws.merge_cells(start_row=subtotal, start_column=1, end_row=subtotal, end_column=3)
    ws.cell(subtotal, 1, "합계 (VAT 별도)")
    ws.cell(subtotal, 4, f"=SUM(D{start}:D{subtotal - 1})")
    ws.cell(subtotal, 5, f"=SUM(E{start}:E{subtotal - 1})")
    ws.cell(subtotal, 6, "")
    ws.merge_cells(start_row=total, start_column=1, end_row=total, end_column=3)
    ws.cell(total, 1, "최종 합계 (VAT 포함)")
    ws.merge_cells(start_row=total, start_column=4, end_row=total, end_column=5)
    ws.cell(total, 4, f"=D{subtotal}+E{subtotal}")
    # Do not retain a static Korean total beside a live editable formula.
    ws.cell(total, 6, "부가가치세 포함")
    ws["C10"] = f"=D{total}"
    for row in (subtotal, total):
        for col in range(1, 7):
            c = ws.cell(row, col); c.border = border; c.font = Font(name="맑은 고딕", size=10, bold=True)
            c.alignment = Alignment(horizontal="center" if col in (1, 6) else "right", vertical="center", wrap_text=True)
            c.fill = PatternFill("solid", fgColor=gray if row == subtotal else "E2F0D9")
    ws.cell(subtotal, 4).number_format = '#,##0"원"'
    ws.cell(subtotal, 5).number_format = '#,##0"원"'
    ws.cell(total, 4).number_format = '#,##0"원"'
    for col, width in enumerate((36, 9, 15, 16, 14, 30), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(1, total + 1):
        ws.row_dimensions[row].height = 24
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 30
    ws.print_area = f"A1:F{total}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&P / &N"
    if seal:
        from app.services.seal import quotation_seal_image
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        stamp = Image(BytesIO(quotation_seal_image(seal)))
        stamp.width = seal['width_mm'] * 96 / 25.4
        stamp.height = seal['height_mm'] * 96 / 25.4
        stamp.anchor = OneCellAnchor(
            _from=AnchorMarker(col=5, colOff=pixels_to_EMU(8), row=5, rowOff=pixels_to_EMU(4)),
            ext=XDRPositiveSize2D(pixels_to_EMU(stamp.width), pixels_to_EMU(stamp.height)),
        )
        ws['E6'] = values.get('representative', '')
        ws.row_dimensions[6].height = seal['height_mm'] * 72 / 25.4 + 8
        ws.add_image(stamp)
    out = BytesIO(); wb.save(out)
    return out.getvalue()


def _q(ns, name):
    return f"{{{ns}}}{name}"


def _hp_paragraph(text: str, *, para_style=23, char_style=16):
    p = etree.Element(_q(HP, "p"), id="0", paraPrIDRef=str(para_style), styleIDRef="0", pageBreak="0", columnBreak="0", merged="0")
    run = etree.SubElement(p, _q(HP, "run"), charPrIDRef=str(char_style))
    t = etree.SubElement(run, _q(HP, "t"))
    for idx, line in enumerate(str(text or "").splitlines() or [""]):
        if idx == 0:
            t.text = line
        else:
            etree.SubElement(t, _q(HP, "lineBreak")).tail = line
    return p


def _hp_table(rows: list[list[str]], widths: list[int], *, header_rows=0, label_cols=0, body_height=2200):
    total_width = sum(widths)
    widths = [round(w * HWP_CONTENT_WIDTH / total_width) for w in widths]
    widths[-1] += HWP_CONTENT_WIDTH - sum(widths)
    table = etree.Element(_q(HP, "tbl"), id="1", zOrder="0", numberingType="TABLE", textWrap="TOP_AND_BOTTOM", textFlow="BOTH_SIDES", lock="0", dropcapstyle="None", pageBreak="CELL", repeatHeader="1", rowCnt=str(len(rows)), colCnt=str(len(widths)), cellSpacing="0", borderFillIDRef="3", noAdjust="0")
    height = header_rows * 2200 + (len(rows) - header_rows) * body_height
    etree.SubElement(table, _q(HP, "sz"), width=str(sum(widths)), widthRelTo="ABSOLUTE", height=str(max(2200, height)), heightRelTo="ABSOLUTE", protect="0")
    etree.SubElement(table, _q(HP, "pos"), treatAsChar="1", affectLSpacing="0", flowWithText="1", allowOverlap="0", holdAnchorAndSO="0", vertRelTo="PARA", horzRelTo="COLUMN", vertAlign="TOP", horzAlign="LEFT", vertOffset="0", horzOffset="0")
    etree.SubElement(table, _q(HP, "outMargin"), left="0", right="0", top="180", bottom="180")
    etree.SubElement(table, _q(HP, "inMargin"), left="300", right="300", top="160", bottom="160")
    for r_idx, row in enumerate(rows):
        tr = etree.SubElement(table, _q(HP, "tr"))
        for c_idx, width in enumerate(widths):
            header = r_idx < header_rows
            label = c_idx < label_cols
            fill = "4" if header or label else "3"
            tc = etree.SubElement(tr, _q(HP, "tc"), name="", header="1" if header else "0", hasMargin="0", protect="0", editable="0", dirty="0", borderFillIDRef=fill)
            sub = etree.SubElement(tc, _q(HP, "subList"), id="", textDirection="HORIZONTAL", lineWrap="BREAK", vertAlign="CENTER", linkListIDRef="0", linkListNextIDRef="0", textWidth="0", textHeight="0", hasTextRef="0", hasNumRef="0")
            text = row[c_idx] if c_idx < len(row) else ""
            sub.append(_hp_paragraph(text, para_style=21 if header or label else 23, char_style=15 if header or label else 16))
            etree.SubElement(tc, _q(HP, "cellAddr"), colAddr=str(c_idx), rowAddr=str(r_idx))
            etree.SubElement(tc, _q(HP, "cellSpan"), colSpan="1", rowSpan="1")
            etree.SubElement(tc, _q(HP, "cellSz"), width=str(width), height=str(2200 if header else body_height))
            etree.SubElement(tc, _q(HP, "cellMargin"), left="300", right="300", top="160", bottom="160")
    p = _hp_paragraph("")
    run = p.find(_q(HP, "run"))
    run.insert(0, table)
    return p


def _hwpx_section(definition: dict, values: dict, seal=None, seal_style=None) -> tuple[bytes, str]:
    values = presentation_values(values)
    with ZipFile(BASE_HWPX) as z:
        root = etree.fromstring(z.read("Contents/section0.xml"))
    page = root.find('.//' + _q(HP, 'pagePr'))
    page.set('width', '59528')
    page.set('height', '84189')
    margin = page.find(_q(HP, 'margin'))
    for key, value in {'left': '6236', 'right': '6236', 'top': '5102', 'bottom': '4819', 'header': '0', 'footer': '0'}.items():
        margin.set(key, value)
    preview = [definition.get("title", "서류")]
    root.append(_hp_paragraph(definition.get("title", "서류"), para_style=21, char_style=13))
    root.append(_hp_paragraph(""))
    for section in definition.get("sections", []):
        kind = section.get("type")
        if kind == "page_break":
            title = _hp_paragraph(section["title"], para_style=21, char_style=13)
            title.set("pageBreak", "1")
            root.append(title)
            preview.append(section["title"])
        elif kind == "signature":
            for text in (values.get("submit_date", ""), f"주 소 : {values.get('address', '')}", f"상 호 : {values.get('company_name', '')}", f"대표자 : {values.get('representative', '')}  (인)", _fill(definition.get("recipient", ""), values)):
                root.append(_hp_paragraph(text, para_style=22))
                preview.append(text)
        elif kind in ("info", "bank"):
            rows = [[label, str(values.get(key, "") or "")] for label, key in section.get("rows", [])]
            root.append(_hp_table(rows, [12500, 37500], label_cols=1))
            preview += [f"{a}: {b}" for a, b in rows]
        elif kind == "paragraph":
            text = _fill(section.get("text", ""), values)
            root.append(_hp_paragraph(text)); preview.append(text)
        elif kind == "field_block":
            label, value = section["label"], str(values.get(section["key"], "") or "")
            root.append(_hp_table([[label], [value]], [50000], header_rows=1)); preview += [label, value]
        elif kind == "line_items":
            items = parse_items(values.get(section.get("key"), ""), values.get("description") or values.get("contract_title", ""), money(values.get("supply_amount")))
            rows = [["항 목", "수량", "단가", "공급가액", "세액", "합계"]]
            supply_total = tax_total = 0
            for item in items:
                supply, tax = item_amounts(item)
                supply_total += supply
                tax_total += tax
                rows.append([item["name"], f"{item['qty']:g}", f"{item['unit']:,}", f"{supply:,}", f"{tax:,}", f"{supply + tax:,}"])
            rows.append(["합 계", "", "", f"{supply_total:,}", f"{tax_total:,}", f"{supply_total + tax_total:,}"])
            root.append(_hp_table(rows, [18000, 4500, 7500, 7500, 5500, 7000], header_rows=1)); preview += [" | ".join(r) for r in rows]
        elif kind == "numbered":
            for idx, item in enumerate(section.get("items", []), 1):
                text = f"{idx}. {_fill(item, values)}"; root.append(_hp_paragraph(text)); preview.append(text)
        elif kind == "attachments":
            root.append(_hp_paragraph("붙임", char_style=15)); preview.append("붙임")
            for idx, item in enumerate(attachment_items(section, values), 1):
                text = f"{idx}. {_fill(item, values)}"; root.append(_hp_paragraph(text)); preview.append(text)
        elif kind in ("stamp_box", "photo_boxes"):
            labels = section.get("labels", [])
            root.append(_hp_table([labels, ["" if kind == "stamp_box" else "[사진 첨부]" for _ in labels]], [25000] * len(labels), header_rows=1, body_height=9100))
            preview += labels
        root.append(_hp_paragraph(""))
    if definition.get("signature"):
        for text in (
            values.get("submit_date") or "    년    월    일",
            f"주 소 : {values.get('address', '')}", f"상 호 : {values.get('company_name', '')}",
            f"대표자 : {values.get('representative', '')}  (인)",
        ):
            root.append(_hp_paragraph(text, para_style=22, char_style=16)); preview.append(text)
    if definition.get("recipient"):
        text = _fill(definition["recipient"], values)
        root.append(_hp_paragraph(text, para_style=23, char_style=14)); preview.append(text)
    if seal:
        from app.services.seal import hwpx_picture
        for index, p in enumerate(root.findall(_q(HP, 'p'))):
            text = ''.join(p.itertext())
            if text.startswith('대표자 :') and '(인)' in text:
                for node in p.findall('.//' + _q(HP, 't')):
                    if node.text and '(인)' in node.text:
                        node.text = node.text.rsplit('(인)', 1)[0].rstrip() + '             (인)'
                p.set('paraPrIDRef', str(seal_style))
                p.find(_q(HP, 'run')).insert(0, hwpx_picture(seal, HWP_CONTENT_WIDTH, 2000 + index))
    for idx, node in enumerate(root.iter(_q(HP, "tbl")), 1):
        node.set("id", str(1000 + idx))
    for idx, node in enumerate(root.iter(_q(HP, "p")), 1):
        node.set("id", str(idx))
    xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return xml, "\n".join(preview)


def render_hwpx(definition: dict, values: dict, seal=None) -> bytes:
    with ZipFile(BASE_HWPX) as src:
        header = etree.fromstring(src.read('Contents/header.xml'))
    seal_style = None
    if seal:
        hh = 'http://www.hancom.co.kr/hwpml/2011/head'
        styles = header.findall('.//' + _q(hh, 'paraPr'))
        seal_style = max(int(node.get('id')) for node in styles) + 1
        style = deepcopy(next(node for node in styles if node.get('id') == '22'))
        style.set('id', str(seal_style))
        width = seal['width_mm'] * 7200 / 25.4
        height = seal['height_mm'] * 7200 / 25.4
        for margin in style.findall('.//' + _q(hh, 'margin')):
            factor = 2 if margin.getparent().tag == _q(HP, 'default') else 1
            for child in margin:
                local = etree.QName(child).localname
                if local == 'right':
                    child.set('value', str(round(max(0, (width - 1750) / 2) * factor)))
                elif local in ('prev', 'next'):
                    child.set('value', str(round((height / 2 + 300) * factor)))
        container = styles[0].getparent()
        container.append(style)
        container.set('itemCnt', str(len(container)))
    section, preview = _hwpx_section(definition, values, seal, seal_style)
    out = BytesIO()
    with ZipFile(BASE_HWPX) as src, ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "Contents/section0.xml":
                data = section
            elif item.filename == "Contents/header.xml":
                for brush in header.xpath('//*[@faceColor]'):
                    if brush.get('faceColor') == '#DCE6F1':
                        brush.set('faceColor', '#F2F3F5')
                data = etree.tostring(header, xml_declaration=True, encoding='UTF-8', standalone=True)
            elif item.filename == 'Contents/content.hpf' and seal:
                package = etree.fromstring(data)
                opf = 'http://www.idpf.org/2007/opf/'
                manifest = package.find(_q(opf, 'manifest'))
                etree.SubElement(manifest, _q(opf, 'item'), id='companySeal', href='BinData/companySeal.png', **{'media-type': 'image/png', 'isEmbeded': '1'})
                data = etree.tostring(package, xml_declaration=True, encoding='UTF-8', standalone=True)
            elif item.filename == "Preview/PrvText.txt":
                data = preview.encode("utf-8")
            info = ZipInfo(item.filename)
            info.compress_type = ZIP_STORED if item.filename == "mimetype" else ZIP_DEFLATED
            dst.writestr(info, data)
        if seal:
            dst.writestr('BinData/companySeal.png', seal['data'], compress_type=ZIP_DEFLATED)
    return out.getvalue()


def render_native(definition: dict, values: dict, seal=None) -> tuple[bytes, str, str]:
    output = definition.get("output_format", "hwpx")
    if output == "xlsx":
        return render_quotation_xlsx(values, seal=seal), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return render_hwpx(definition, values, seal=seal), "hwpx", "application/hwp+zip"
