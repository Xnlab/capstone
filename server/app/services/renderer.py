"""서류 파일 생성: 기본 양식/자체 양식 → DOCX → (PDF, HTML 미리보기).

모든 서류를 먼저 DOCX로 만들고, PDF와 미리보기는 그 DOCX 하나에서 만든다(출력 경로 단일화).
"""

import html
import re
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.core.config import settings
from app.standard_fields import field_label

FONT = "맑은 고딕"
_PH = re.compile(r"\{(\w+)\}")


def _fill(text: str, values: dict) -> str:
    return _PH.sub(lambda m: str(values.get(m.group(1)) or f"[{field_label(m.group(1))}]"), text)


def _set_font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _shade(cell, hex_color="EFEFF4"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _para(doc, text="", size=11, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if text:
        _set_font(p.add_run(text), size, bold)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


# ---------------------------------------------------------------- 기본 양식
def render_builtin_docx(definition: dict, values: dict) -> DocxDocument:
    doc = DocxDocument()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.5)

    _para(doc, definition.get("title", ""), size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    rows = definition.get("table", [])
    if rows:
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, key) in enumerate(rows):
            lc, vc = table.rows[i].cells
            lc.width, vc.width = Cm(4), Cm(12)
            lc.text, vc.text = "", ""
            _set_font(lc.paragraphs[0].add_run(label), 11, True)
            lc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _shade(lc)
            _set_font(vc.paragraphs[0].add_run(str(values.get(key, "") or "")), 11)
        _para(doc, space_after=12)

    for line in _fill(definition.get("body", ""), values).split("\n"):
        _para(doc, line, size=11, space_after=4)

    atts = definition.get("attachments") or []
    if atts:
        _para(doc, space_after=6)
        _para(doc, "붙임", size=11, bold=True, space_after=2)
        for i, a in enumerate(atts, 1):
            _para(doc, f"  {i}. {a}", size=11, space_after=2)

    if definition.get("signature"):
        _para(doc, space_after=18)
        _para(doc, values.get("submit_date") or "    년    월    일", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
        for lab, key in (("상    호", "company_name"), ("주    소", "address"), ("대 표 자", "representative")):
            suffix = "   (인)" if key == "representative" else ""
            _para(doc, f"{lab} : {values.get(key, '') or ''}{suffix}", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)

    if definition.get("recipient"):
        _para(doc, space_after=18)
        _para(doc, _fill(definition["recipient"], values), size=14, bold=True, space_after=0)
    return doc


# ---------------------------------------------------------------- 기업 자체 양식
def _write_cell(cell, value: str):
    first = cell.paragraphs[0]
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    _set_font(first.add_run(value))


def _write_after_colon(paragraph, value: str):
    text = paragraph.text
    idx = max(text.rfind(":"), text.rfind("："))
    head = text[: idx + 1] if idx >= 0 else text
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    _set_font(paragraph.add_run(f"{head} {value}"))


def fill_uploaded_docx(source_path: Path, fields: list[dict], values: dict) -> DocxDocument:
    """원본 양식을 그대로 두고, 분석된 입력 위치에만 값을 채운다."""
    doc = DocxDocument(str(source_path))
    for f in fields:
        key, loc = f.get("field_key"), f.get("location") or {}
        if not key:
            continue
        val = str(values.get(key, "") or "")
        if not val:
            continue
        if loc.get("type") == "cell":
            try:
                cell = doc.tables[loc["table"]].rows[loc["row"]].cells[loc["col"]]
                _write_cell(cell, val)
            except IndexError:
                continue
        elif loc.get("type") == "para":
            if loc["index"] < len(doc.paragraphs):
                _write_after_colon(doc.paragraphs[loc["index"]], val)
    return doc


def rebuild_from_pdf_fields(title: str, fields: list[dict], values: dict) -> DocxDocument:
    """PDF 양식은 원본 편집 대신 분석된 항목 순서대로 Word 문서를 재구성한다."""
    rows = [[f["label"], f["field_key"]] for f in fields if f.get("field_key")]
    return render_builtin_docx({"title": title or "서류", "table": rows, "signature": True}, values)


def save_docx(doc: DocxDocument, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def docx_bytes(doc: DocxDocument) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- DOCX → 블록(문단/표)
def _iter_blocks(doc: DocxDocument):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, doc)


def _para_style(p: DocxParagraph) -> dict:
    size, bold = 11, False
    for r in p.runs:
        if r.font.size:
            size = r.font.size.pt
        bold = bold or bool(r.bold)
    align = {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right"}.get(p.alignment, "left")
    return {"size": size, "bold": bold, "align": align}


def _table_rows(t: DocxTable) -> list[list[str]]:
    rows = []
    for row in t.rows:
        seen, cells = set(), []
        for c in row.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            cells.append(c.text.strip())
        rows.append(cells)
    return rows


def docx_to_html(doc: DocxDocument) -> str:
    out = []
    for b in _iter_blocks(doc):
        if isinstance(b, DocxParagraph):
            st = _para_style(b)
            txt = html.escape(b.text) or "&nbsp;"
            weight = "700" if st["bold"] else "400"
            out.append(f'<p style="text-align:{st["align"]};font-size:{st["size"]}pt;font-weight:{weight};margin:0 0 4px">{txt}</p>')
        else:
            rows = _table_rows(b)
            trs = "".join(
                "<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in rows
            )
            out.append(f'<table class="doc-table">{trs}</table>')
    return "\n".join(out)


def docx_to_text(doc: DocxDocument) -> str:
    parts = []
    for b in _iter_blocks(doc):
        if isinstance(b, DocxParagraph):
            parts.append(b.text)
        else:
            parts += [" | ".join(r) for r in _table_rows(b)]
    return "\n".join(p for p in parts if p.strip())


# ---------------------------------------------------------------- DOCX → PDF (reportlab)
_font_registered = False


def _register_font() -> str:
    global _font_registered
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if _font_registered:
        return "KFont"
    path = settings.PDF_FONT_PATH
    if Path(path).exists():
        pdfmetrics.registerFont(TTFont("KFont", path))
    else:
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
        return "HYSMyeongJo-Medium"
    _font_registered = True
    return "KFont"


def docx_to_pdf_bytes(doc: DocxDocument) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_font()
    align_map = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT}
    story = []
    width = A4[0] - 5 * cm
    for b in _iter_blocks(doc):
        if isinstance(b, DocxParagraph):
            st = _para_style(b)
            if not b.text.strip():
                story.append(Spacer(1, 8))
                continue
            style = ParagraphStyle(
                "p", fontName=font, fontSize=st["size"], leading=st["size"] * 1.5,
                alignment=align_map[st["align"]], spaceAfter=4,
            )
            text = html.escape(b.text).replace("\n", "<br/>")
            story.append(Paragraph(f"<b>{text}</b>" if st["bold"] and font == "KFont" else text, style))
        else:
            rows = _table_rows(b)
            if not rows:
                continue
            cell_style = ParagraphStyle("c", fontName=font, fontSize=10, leading=14)
            data = [[Paragraph(html.escape(c).replace("\n", "<br/>"), cell_style) for c in r] for r in rows]
            ncols = max(len(r) for r in data)
            data = [r + [""] * (ncols - len(r)) for r in data]
            col_w = [width * 0.28] + [width * 0.72 / (ncols - 1)] * (ncols - 1) if ncols > 1 else [width]
            t = Table(data, colWidths=col_w)
            ts = [("GRID", (0, 0), (-1, -1), 0.6, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
            if ncols > 1:
                ts.append(("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFEFF4")))
            t.setStyle(TableStyle(ts))
            story += [t, Spacer(1, 10)]
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2.5 * cm, bottomMargin=2 * cm).build(story)
    return buf.getvalue()

