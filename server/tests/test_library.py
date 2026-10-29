import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as Word
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.database import Base, get_db
from app.core.security import get_current_user
from app.models import Contract, Document, Template, User, LibraryAdmin
from app.routers import contracts, templates


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.admin = User(email="admin@example.test", password_hash="unused")
        self.user = User(email="member@example.test", password_hash="unused")
        self.db.add_all([self.admin, self.user]); self.db.commit()
        self.db.add(LibraryAdmin(user_id=self.admin.id)); self.db.commit()
        self.current = self.admin
        self.patch = patch.object(settings, "LIBRARY_ADMIN_EMAILS", {self.admin.email})
        self.patch.start()
        app = FastAPI()
        app.include_router(templates.router)
        app.include_router(contracts.router)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.current
        self.client = TestClient(app)
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.client.close(); self.db.close(); self.engine.dispose(); self.patch.stop(); self.tmp.cleanup()

    def form(self, **overrides):
        path = Path(self.tmp.name) / "form.docx"
        doc = Word(); doc.add_paragraph("지정 양식 고정 문구"); doc.add_paragraph("상호 : ____"); doc.save(path)
        info = dict(scope="organization", organization="주식회사 테스트", document_type="착수계", source_url="https://example.test/form", version="2026", usage_terms="검토 완료", review_note="샘플 출력 확인")
        info.update(overrides)
        t = Template(owner_id=self.admin.id, kind="uploaded", name="착수계", aliases=["착수계"], definition={"library": info}, fields=[{"label": "상호", "field_key": "company_name", "location": {"type": "para", "index": 1}}], source_path=str(path), source_type="docx", status="review_pending")
        self.db.add(t); self.db.commit()
        return t

    def publish(self, t):
        return self.client.post(f"/templates/{t.id}/publish", json=dict(mapping_checked=True, output_checked=True, rights_checked=True))

    def test_admin_required_and_drafts_private(self):
        t = self.form(); self.current = self.user
        self.assertEqual(self.publish(t).status_code, 403)
        self.assertEqual(self.client.get(f"/templates/{t.id}").status_code, 404)
        self.assertEqual(self.client.get("/templates?include_drafts=true").json(), [])

    def test_publish_and_protect_public_form(self):
        t = self.form(); self.assertEqual(self.publish(t).status_code, 200)
        self.current = self.user
        self.assertEqual(len(self.client.get("/templates").json()), 1)
        self.assertEqual(self.client.put(f"/templates/{t.id}/mapping", json={"fields": []}).status_code, 403)
        self.assertEqual(self.client.delete(f"/templates/{t.id}").status_code, 403)

    def test_pdf_and_unmapped_forms_cannot_publish(self):
        t = self.form(); t.source_type = "pdf"; self.db.commit()
        self.assertEqual(self.publish(t).status_code, 400)
        t.source_type = "docx"; t.fields = []; self.db.commit()
        self.assertEqual(self.publish(t).status_code, 400)

    def test_issuer_and_project_matching(self):
        t = self.form(project_code="BID-1"); self.publish(t); self.current = self.user
        self.assertEqual(self.client.get("/templates/search", params={"client_name": "테스트(주)"}).json()["results"], [])
        results = self.client.get("/templates/search", params={"client_name": "테스트(주)", "project_code": "BID-1"}).json()["results"]
        self.assertEqual(results[0]["template"]["id"], t.id)
        self.assertEqual(self.client.get("/templates/search", params={"client_name": "다른회사", "project_code": "BID-1"}).json()["results"], [])

    def test_contract_guard_and_version_immutability(self):
        t = self.form(); self.publish(t)
        c = Contract(user_id=self.user.id, contract_title="테스트", client_name="다른회사")
        self.db.add(c); self.db.commit(); self.current = self.user
        self.assertEqual(self.client.post(f"/contracts/{c.id}/documents", json={"template_id": t.id}).status_code, 400)
        c.client_name = "테스트(주)"; self.db.commit()
        self.assertEqual(self.client.post(f"/contracts/{c.id}/documents", json={"template_id": t.id}).status_code, 200)
        self.current = self.admin
        self.assertEqual(self.client.put(f"/templates/{t.id}/mapping", json={"fields": []}).status_code, 409)

    def test_original_and_sample_preserve_fixed_text(self):
        from io import BytesIO
        t = self.form()
        sample = self.client.get(f"/templates/{t.id}/sample-file")
        self.assertEqual(sample.status_code, 200)
        text = "\n".join(p.text for p in Word(BytesIO(sample.content)).paragraphs)
        self.assertIn("지정 양식 고정 문구", text)
        self.assertIn("주식회사 샘플", text)
        self.assertEqual(self.client.get(f"/templates/{t.id}/original").content, Path(t.source_path).read_bytes())

    def test_private_upload_not_visible_even_to_library_admin(self):
        t = self.form(scope="private"); t.owner_id = self.user.id; t.status = "ready"; self.db.commit()
        self.assertEqual(self.client.get(f"/templates/{t.id}").status_code, 404)

    def test_favorites_are_idempotent_and_user_scoped(self):
        t = self.form(); self.publish(t)
        for _ in range(2):
            self.assertEqual(self.client.put(f"/templates/{t.id}/favorite", json={"favorite": True}).status_code, 200)
        self.assertEqual(self.client.get("/templates/favorites").json(), [t.id])
        self.current = self.user
        self.assertEqual(self.client.get("/templates/favorites").json(), [])
        self.current = self.admin
        self.client.put(f"/templates/{t.id}/favorite", json={"favorite": False})
        self.assertEqual(self.client.get("/templates/favorites").json(), [])

    def test_rename_preserves_document_title_and_checks_permissions(self):
        t = self.form(); self.publish(t)
        c = Contract(user_id=self.user.id, contract_title="계약", client_name="테스트")
        self.db.add(c); self.db.flush()
        d = Document(contract_id=c.id, template_id=t.id, title=t.name, values={})
        self.db.add(d); self.db.commit()
        self.assertEqual(self.client.put(f"/templates/{t.id}/name", json={"name": "새 양식 이름"}).status_code, 200)
        self.db.refresh(d); self.assertEqual(d.title, "착수계")
        self.assertEqual(self.client.put(f"/templates/{t.id}/name", json={"name": "   "}).status_code, 422)
        self.current = self.user
        self.assertEqual(self.client.put(f"/templates/{t.id}/name", json={"name": "변경"}).status_code, 403)

    def test_organization_suggestions_do_not_leak_other_contracts(self):
        t = self.form(); self.publish(t)
        self.db.add_all([Contract(user_id=self.user.id, contract_title="개인", client_name="비공개 거래처"), Contract(user_id=self.admin.id, contract_title="본인", client_name="이전 거래처")]); self.db.commit()
        names = self.client.get("/templates/organizations").json()
        self.assertIn("주식회사 테스트", names)
        self.assertIn("이전 거래처", names)
        self.assertNotIn("비공개 거래처", names)

    def pdf_form(self):
        import pymupdf
        t = self.form()
        path = Path(self.tmp.name) / "original.pdf"
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=400, height=500)
            page.insert_text((50, 40), "FIXED ORIGINAL CLAUSE")
            page.insert_text((50, 200), "2025")
            pdf.save(path)
        t.source_type, t.source_path = "pdf", str(path)
        t.fields = [{"label": "회사명", "field_key": "company_name", "location": {"type": "pdf_rect", "page": 0, "rect": [50, 80, 350, 110]}}]
        self.db.commit()
        return t

    def test_positioned_pdf_download_preserves_original(self):
        import pymupdf
        t = self.pdf_form()
        original = Path(t.source_path).read_bytes()
        self.assertEqual(self.publish(t).status_code, 200)
        response = self.client.get(f"/templates/{t.id}/sample-file")
        self.assertEqual(response.headers["content-type"], "application/pdf")
        with pymupdf.open(stream=response.content, filetype="pdf") as pdf:
            self.assertIn("FIXED ORIGINAL CLAUSE", pdf[0].get_text())
            self.assertIn("주식회사 샘플", pdf[0].get_text())
        self.assertEqual(original, Path(t.source_path).read_bytes())

    def test_pdf_rejects_fixed_text_overlap_and_overflow(self):
        from app.services.pdf_forms import fill_pdf, validate_locations
        from fastapi import HTTPException
        t = self.pdf_form()
        fields = t.fields
        t.fields = [{**fields[0], "location": {"type": "pdf_rect", "page": 0, "rect": [45, 25, 300, 45]}}]
        with self.assertRaises(HTTPException):
            validate_locations(t)
        t.fields = fields
        with self.assertRaises(HTTPException):
            fill_pdf(t, {"company_name": "너무 긴 회사명" * 1000})

    def test_pdf_cannot_erase_non_date_or_use_bad_coordinates(self):
        from app.services.pdf_forms import validate_locations
        from fastapi import HTTPException
        t = self.pdf_form()
        t.fields = [{"label": "회사", "field_key": "company_name", "location": {"type": "pdf_rect", "page": 0, "rect": [50, 185, 100, 210], "replace": True}}]
        with self.assertRaises(HTTPException): validate_locations(t)
        t.fields = [{"label": "회사", "field_key": "company_name", "location": {"type": "pdf_rect", "page": 0, "rect": [-50, 80, 300, 110]}}]
        with self.assertRaises(HTTPException): validate_locations(t)


if __name__ == "__main__":
    unittest.main()
